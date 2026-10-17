"""Build the shared tester checklist as an .xlsx made for Google Sheets.

Usage: python build_sheet.py <checklist.json> <out.xlsx> <version>

checklist.json: {"sections": [{"title", "items": [{"id", "title", "client", "important", "try", "expect"}]}]}
(client is "Both", "Retail" or "WoW Forever"). Every ship of the patch notes comes with one of these sheets, saved
next to the tester zips. One shared sheet for every tester: each tester takes a Result + Note column pair (rename
"Tester 1" to your name), Result is a Good / Not good dropdown coloured green / red. Upload to Google Drive, open
with Google Sheets, save as a Google Sheet, share with edit rights: the dropdowns, colours, frozen panes and merged
section rows carry over.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if len(sys.argv) != 4:
    sys.exit(__doc__)
src, out, version = sys.argv[1], sys.argv[2], sys.argv[3]
TESTERS = 6

data = json.load(open(src, encoding="utf-8"))
sections = data["sections"]

BASE = ["#", "Check", "Client", "Priority", "What to try", "What should happen"]
BASE_W = [7, 28, 13, 11, 60, 46]
RESULT_W, NOTE_W = 13, 30

DARK = PatternFill("solid", fgColor="263238")
TESTER_HEAD = PatternFill("solid", fgColor="37474F")
SECTION = PatternFill("solid", fgColor="D7F5EC")
RESULT_FILL = PatternFill("solid", fgColor="FFF8E1")
GOOD_FILL = PatternFill("solid", fgColor="C8E6C9")
BAD_FILL = PatternFill("solid", fgColor="FFCDD2")
THIN = Side(style="thin", color="BDBDBD")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
CLIENT_COLOR = {"Both": "37474F", "Retail": "1565C0", "WoW Forever": "E65100"}

ncol = len(BASE) + 2 * TESTERS
wb = Workbook()
ws = wb.active
ws.title = "Checklist"
for i, w in enumerate(BASE_W, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for t in range(TESTERS):
    ws.column_dimensions[get_column_letter(len(BASE) + 1 + 2 * t)].width = RESULT_W
    ws.column_dimensions[get_column_letter(len(BASE) + 2 + 2 * t)].width = NOTE_W

ws["A1"] = "EllesmereUI %s test checklist" % version
ws["A1"].font = Font(bold=True, size=15)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(BASE))
ws["A2"] = ("Write your name over one Tester pair (row 4). For each check you can do, pick Good or Not good in your "
            "Result column; for anything Not good, say what happened in your Note column (paste any Lua error, add "
            "your client and class). Leave a check blank if you could not try it. Client says which game version a "
            "check is for; Important checks matter most.")
ws["A2"].font = Font(italic=True, color="455A64")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(BASE))
ws.row_dimensions[2].height = 48

NAME_ROW, HEADER_ROW = 4, 5
for c, name in enumerate(BASE, 1):
    cell = ws.cell(row=HEADER_ROW, column=c, value=name)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = DARK
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = GRID
for t in range(TESTERS):
    c0 = len(BASE) + 1 + 2 * t
    ws.cell(row=NAME_ROW, column=c0, value="Tester %d" % (t + 1))
    ws.merge_cells(start_row=NAME_ROW, start_column=c0, end_row=NAME_ROW, end_column=c0 + 1)
    name = ws.cell(row=NAME_ROW, column=c0)
    name.font = Font(bold=True, color="FFFFFF")
    name.fill = TESTER_HEAD
    name.alignment = Alignment(horizontal="center", vertical="center")
    for k, label in enumerate(("Result", "Note")):
        cell = ws.cell(row=HEADER_ROW, column=c0 + k, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID
ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=3)

dv = DataValidation(type="list", formula1='"Good,Not good"', allow_blank=True,
                    errorTitle="Good or Not good", error="Pick Good or Not good from the list.", showErrorMessage=True)
ws.add_data_validation(dv)

r = HEADER_ROW + 1
first = r
count = 0
for sec in sections:
    ws.cell(row=r, column=1, value=sec["title"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(BASE))
    head = ws.cell(row=r, column=1)
    head.font = Font(bold=True, size=12, color="004D40")
    head.fill = SECTION
    head.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 22
    r += 1
    for it in sec["items"]:
        client = it.get("client", "Both")
        values = [it.get("id", ""), it["title"], client, "Important" if it.get("important") else "",
                  it["try"], it["expect"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = WRAP_TOP
            cell.border = GRID
        ws.cell(row=r, column=1).alignment = CENTER_TOP
        ws.cell(row=r, column=2).font = Font(bold=True)
        cc = ws.cell(row=r, column=3)
        cc.alignment = CENTER_TOP
        cc.font = Font(bold=True, color=CLIENT_COLOR.get(client, "37474F"))
        pc = ws.cell(row=r, column=4)
        pc.alignment = CENTER_TOP
        pc.font = Font(bold=True, color="B71C1C")
        for t in range(TESTERS):
            c0 = len(BASE) + 1 + 2 * t
            res = ws.cell(row=r, column=c0)
            res.fill = RESULT_FILL
            res.border = GRID
            res.alignment = Alignment(horizontal="center", vertical="top")
            dv.add(res)
            note = ws.cell(row=r, column=c0 + 1)
            note.border = GRID
            note.alignment = WRAP_TOP
        r += 1
        count += 1
last = r - 1

for t in range(TESTERS):
    col = get_column_letter(len(BASE) + 1 + 2 * t)
    rng = "%s%d:%s%d" % (col, first, col, last)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Good"'], fill=GOOD_FILL,
                                                  font=Font(bold=True, color="1B5E20")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not good"'], fill=BAD_FILL,
                                                  font=Font(bold=True, color="B71C1C")))

wb.save(out)
print("saved", out, "|", len(sections), "sections,", count, "checks,", TESTERS, "tester column pairs")
